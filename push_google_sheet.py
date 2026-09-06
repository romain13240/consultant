import os,json,sys
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import load_workbook

HOME=os.path.expanduser('~'); token=json.load(open(os.path.join(HOME,'.hermes/google_token.json')))
creds=Credentials(token=token.get('token') or token.get('access_token'),refresh_token=token['refresh_token'],token_uri=token['token_uri'],client_id=token['client_id'],client_secret=token['client_secret'],scopes=token.get('scopes'))
drive=build('drive','v3',credentials=creds); sheets=build('sheets','v4',credentials=creds)
# Find or create target folder
q="name = '#Consultant + MRR' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
res=drive.files().list(q=q,spaces='drive',fields='files(id,name)',pageSize=10).execute().get('files',[])
if res: folder=res[0]
else: folder=drive.files().create(body={'name':'#Consultant + MRR','mimeType':'application/vnd.google-apps.folder'},fields='id,name').execute()
# Create native Google Sheet
ss=sheets.spreadsheets().create(body={'properties':{'title':'Consultant agent IA v2'}}).execute(); sid=ss['spreadsheetId']
# Move into folder
file=drive.files().get(fileId=sid,fields='parents').execute();
if file.get('parents'): drive.files().update(fileId=sid,addParents=folder['id'],removeParents=','.join(file['parents']),fields='id,parents').execute()
# Rename default sheet and add others
names=['Synthèse','Paramètres','MRR mensuel','Synthèse annuelle','Temps & capacité']
meta=sheets.spreadsheets().get(spreadsheetId=sid).execute(); first=meta['sheets'][0]['properties']; requests=[{'updateSheetProperties':{'properties':{'sheetId':first['sheetId'],'title':'Synthèse'},'fields':'title'}}]
for n in names[1:]: requests.append({'addSheet':{'properties':{'title':n}}})
sheets.spreadsheets().batchUpdate(spreadsheetId=sid,body={'requests':requests}).execute()
# Load workbook and push values, respecting Google limits
wb=load_workbook('/home/hermes/consultant/Consultant agent IA v2.xlsx',data_only=False)
def clean(v):
    if v is None:return ''
    if isinstance(v,(int,float,str,bool)):return v
    return str(v)
value_data=[]
for n in names:
    ws=wb[n]; values=[[clean(c.value) for c in row] for row in ws.iter_rows()]
    # trim empty trailing rows/cols
    while values and all(x=='' for x in values[-1]): values.pop()
    while values and values[0] and all(row[-1]=='' for row in values):
        for row in values: row.pop()
    value_data.append({'range':f"'{n}'!A1",'values':values})
sheets.spreadsheets().values().batchUpdate(spreadsheetId=sid,body={'valueInputOption':'USER_ENTERED','data':value_data}).execute()
# Basic formatting
meta=sheets.spreadsheets().get(spreadsheetId=sid).execute(); idmap={s['properties']['title']:s['properties']['sheetId'] for s in meta['sheets']}
req=[]
for n in names:
    sid2=idmap[n]; req.append({'repeatCell':{'range':{'sheetId':sid2,'startRowIndex':0,'endRowIndex':1},'cell':{'userEnteredFormat':{'textFormat':{'bold':True,'foregroundColor':{'red':1,'green':1,'blue':1}},'backgroundColor':{'red':0,'green':0.443,'blue':0.89}}},'fields':'userEnteredFormat(textFormat,backgroundColor)'}})
    req.append({'updateSheetProperties':{'properties':{'sheetId':sid2,'gridProperties':{'frozenRowCount':3}},'fields':'gridProperties.frozenRowCount'}})
sheets.spreadsheets().batchUpdate(spreadsheetId=sid,body={'requests':req}).execute()
print(json.dumps({'folder_id':folder['id'],'spreadsheet_id':sid,'url':f'https://docs.google.com/spreadsheets/d/{sid}/edit','sheets':names},ensure_ascii=False))
