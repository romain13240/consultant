#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
push_to_drive.py — Publie le classeur « Consultant agent IA v2 » sur Google Drive.

Reprend le classeur produit par build_gsheet.py et le pousse dans un Google Sheet
**natif** : les formules restent vivantes (valueInputOption USER_ENTERED), les
formats de nombre sont conservés, et des graphiques Google natifs sont ajoutés.

Prérequis :
  pip install google-api-python-client google-auth openpyxl
  Un jeton OAuth utilisateur avec les portées Drive + Sheets, par défaut dans
  ~/.hermes/google_token.json (surchargeable par la variable GOOGLE_TOKEN).

Usage :
  python sheets/build_gsheet.py          # (re)générer le classeur
  python sheets/push_to_drive.py         # publier / mettre à jour
  python sheets/push_to_drive.py --new   # forcer la création d'un nouveau fichier
"""

from __future__ import print_function

import json
import os
import sys

try:
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
except ImportError:  # pragma: no cover
    sys.exit(u"""Dépendances Google absentes.

Sur Raspberry Pi OS récent, pip refuse d'installer dans le système
(PEP 668 / externally-managed-environment). Il faut un environnement virtuel :

  sudo apt install -y python3-venv
  python3 -m venv ~/.venvs/consultant
  ~/.venvs/consultant/bin/pip install -q google-api-python-client google-auth openpyxl
  ~/.venvs/consultant/bin/python ~/consultant/sheets/push_to_drive.py

Ou, en une seule commande depuis ~/consultant :  ./deploy-rpi.sh --drive
""")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
ICI = os.path.dirname(os.path.abspath(__file__))
CLASSEUR = os.environ.get("CLASSEUR", os.path.join(ICI, "consultant-agent-ia-v2.xlsx"))
# Le jeton n'est pas au même endroit selon la machine et l'utilisateur : on
# essaie les emplacements connus plutôt que d'en coder un seul en dur.
CANDIDATS_JETON = [
    os.environ.get("GOOGLE_TOKEN"),
    os.path.join(os.path.expanduser("~"), ".hermes", "google_token.json"),
    os.path.join(os.path.expanduser("~"), ".config", "consultant", "google_token.json"),
    os.path.join(os.path.expanduser("~"), ".google_token.json"),
    "/home/hermes/.hermes/google_token.json",
]


def trouver_jeton():
    for c in CANDIDATS_JETON:
        if c and os.path.exists(c):
            return c
    return None
DOSSIER = os.environ.get("DRIVE_FOLDER", u"#Consultant + MRR")
TITRE = os.environ.get("SHEET_TITLE", u"Consultant agent IA v2")

BLEU = {"red": 0.0, "green": 0.443, "blue": 0.890}
CREME = {"red": 1.0, "green": 0.973, "blue": 0.906}
GRIS = {"red": 0.961, "green": 0.961, "blue": 0.969}
BLANC = {"red": 1.0, "green": 1.0, "blue": 1.0}


# --------------------------------------------------------------------------- #
def services():
    jeton = trouver_jeton()
    if not jeton:
        print(u"Jeton OAuth Google introuvable. Emplacements essayés :")
        for c in CANDIDATS_JETON:
            if c:
                print(u"  - %s" % c)
        print(u"")
        print(u"Localisez-le :")
        print(u"  sudo find /home -name 'google_token.json' 2>/dev/null")
        print(u"puis relancez en indiquant le chemin trouvé :")
        print(u"  GOOGLE_TOKEN=/chemin/vers/google_token.json python3 sheets/push_to_drive.py")
        sys.exit(1)
    print(u"Jeton : %s" % jeton)
    tk = json.load(open(jeton))
    creds = Credentials(
        token=tk.get("token") or tk.get("access_token"),
        refresh_token=tk["refresh_token"],
        token_uri=tk["token_uri"],
        client_id=tk["client_id"],
        client_secret=tk["client_secret"],
        scopes=tk.get("scopes"),
    )
    return build("drive", "v3", credentials=creds), build("sheets", "v4", credentials=creds)


def dossier_cible(drive):
    q = (u"name = '%s' and mimeType = 'application/vnd.google-apps.folder' "
         u"and trashed = false" % DOSSIER.replace("'", "\\'"))
    trouves = drive.files().list(q=q, spaces="drive", fields="files(id,name)",
                                 pageSize=10).execute().get("files", [])
    if trouves:
        print(u"Dossier existant : %s (%s)" % (DOSSIER, trouves[0]["id"]))
        return trouves[0]["id"]
    cree = drive.files().create(
        body={"name": DOSSIER, "mimeType": "application/vnd.google-apps.folder"},
        fields="id").execute()
    print(u"Dossier créé : %s (%s)" % (DOSSIER, cree["id"]))
    return cree["id"]


def classeur_cible(drive, sheets_api, folder_id, forcer_nouveau):
    if not forcer_nouveau:
        q = (u"name = '%s' and '%s' in parents and trashed = false"
             % (TITRE.replace("'", "\\'"), folder_id))
        trouves = drive.files().list(q=q, spaces="drive", fields="files(id,name)",
                                     pageSize=5).execute().get("files", [])
        if trouves:
            print(u"Classeur existant réutilisé : %s" % trouves[0]["id"])
            return trouves[0]["id"], False
    ss = sheets_api.spreadsheets().create(body={"properties": {"title": TITRE}},
                                          fields="spreadsheetId").execute()
    sid = ss["spreadsheetId"]
    meta = drive.files().get(fileId=sid, fields="parents").execute()
    drive.files().update(fileId=sid, addParents=folder_id,
                         removeParents=",".join(meta.get("parents", [])),
                         fields="id").execute()
    print(u"Classeur créé : %s" % sid)
    return sid, True


# --------------------------------------------------------------------------- #
def lire_classeur():
    if not os.path.exists(CLASSEUR):
        sys.exit(u"Classeur introuvable : %s\nLancez d'abord build_gsheet.py." % CLASSEUR)
    wb = load_workbook(CLASSEUR, data_only=False)
    onglets = []
    for nom in wb.sheetnames:
        ws = wb[nom]
        lignes, formats, largeurs = [], [], []
        for row in ws.iter_rows():
            ligne, fligne = [], []
            for c in row:
                v = c.value
                if v is None:
                    ligne.append("")
                elif isinstance(v, bool):
                    ligne.append(v)
                elif isinstance(v, (int, float)):
                    ligne.append(v)
                else:
                    ligne.append(u"%s" % v)
                fligne.append(c.number_format or "General")
            lignes.append(ligne)
            formats.append(fligne)
        while lignes and all(x == "" for x in lignes[-1]):
            lignes.pop()
            formats.pop()
        ncols = max([len(l) for l in lignes] or [1])
        for i in range(1, ncols + 1):
            dim = ws.column_dimensions.get(get_column_letter(i))
            largeurs.append(int((dim.width or 10) * 7.5) if dim and dim.width else 90)
        onglets.append({"nom": nom, "lignes": lignes, "formats": formats,
                        "largeurs": largeurs, "nlignes": len(lignes), "ncols": ncols})
    return onglets


def preparer_onglets(sheets_api, sid, onglets, neuf):
    meta = sheets_api.spreadsheets().get(spreadsheetId=sid).execute()
    existants = {s["properties"]["title"]: s["properties"] for s in meta["sheets"]}
    req = []

    # renommer la feuille par défaut du classeur neuf
    if neuf:
        premier = meta["sheets"][0]["properties"]
        req.append({"updateSheetProperties": {
            "properties": {"sheetId": premier["sheetId"], "title": onglets[0]["nom"],
                           "index": 0},
            "fields": "title,index"}})
        existants = {onglets[0]["nom"]: dict(premier, title=onglets[0]["nom"])}

    for i, o in enumerate(onglets):
        if o["nom"] in existants:
            req.append({"updateSheetProperties": {
                "properties": {"sheetId": existants[o["nom"]]["sheetId"], "index": i},
                "fields": "index"}})
        else:
            req.append({"addSheet": {"properties": {
                "title": o["nom"], "index": i,
                "gridProperties": {"rowCount": max(o["nlignes"] + 20, 60),
                                   "columnCount": max(o["ncols"] + 6, 20)}}}})
    if req:
        sheets_api.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": req}).execute()

    meta = sheets_api.spreadsheets().get(spreadsheetId=sid).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}

    # supprimer les onglets qui ne font plus partie du modèle
    a_supprimer = [sid_ for titre, sid_ in ids.items()
                   if titre not in [o["nom"] for o in onglets]]
    if a_supprimer and len(ids) > len(a_supprimer):
        sheets_api.spreadsheets().batchUpdate(
            spreadsheetId=sid,
            body={"requests": [{"deleteSheet": {"sheetId": s}} for s in a_supprimer]}).execute()
        meta = sheets_api.spreadsheets().get(spreadsheetId=sid).execute()
        ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}

    # purger les graphiques déjà présents (ils seront recréés)
    charts = []
    for s in meta["sheets"]:
        for c in s.get("charts", []):
            charts.append({"deleteEmbeddedObject": {"objectId": c["chartId"]}})
    if charts:
        sheets_api.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": charts}).execute()
    return ids


def pousser_valeurs(sheets_api, sid, onglets):
    sheets_api.spreadsheets().values().batchClear(
        spreadsheetId=sid,
        body={"ranges": [u"'%s'" % o["nom"] for o in onglets]}).execute()
    data = [{"range": u"'%s'!A1" % o["nom"], "values": o["lignes"]}
            for o in onglets if o["lignes"]]
    sheets_api.spreadsheets().values().batchUpdate(
        spreadsheetId=sid,
        body={"valueInputOption": "USER_ENTERED", "data": data}).execute()
    print(u"Valeurs et formules poussées (%d onglets)" % len(data))


def _fmt_google(fmt):
    """Traduit un format openpyxl en numberFormat Google Sheets."""
    if not fmt or fmt == "General":
        return None
    f = fmt.replace("\\", "")
    if "%" in f:
        return {"type": "PERCENT", "pattern": f}
    if "mm" in f and "yy" in f:
        return {"type": "DATE", "pattern": f}
    return {"type": "NUMBER", "pattern": f}


def mettre_en_forme(sheets_api, sid, ids, onglets):
    req = []
    for o in onglets:
        gid = ids[o["nom"]]
        # largeurs de colonnes
        for i, w in enumerate(o["largeurs"]):
            req.append({"updateDimensionProperties": {
                "range": {"sheetId": gid, "dimension": "COLUMNS",
                          "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": max(60, min(520, w))},
                "fields": "pixelSize"}})
        # formats de nombre, colonne par colonne (le format dominant l'emporte)
        for col in range(o["ncols"]):
            compte = {}
            for r in range(len(o["formats"])):
                if col < len(o["formats"][r]):
                    f = o["formats"][r][col]
                    if f and f != "General":
                        compte[f] = compte.get(f, 0) + 1
            if not compte:
                continue
            dominant = max(compte.items(), key=lambda kv: kv[1])[0]
            nf = _fmt_google(dominant)
            if not nf:
                continue
            req.append({"repeatCell": {
                "range": {"sheetId": gid, "startColumnIndex": col, "endColumnIndex": col + 1},
                "cell": {"userEnteredFormat": {"numberFormat": nf}},
                "fields": "userEnteredFormat.numberFormat"}})
        # police et fond général
        req.append({"repeatCell": {
            "range": {"sheetId": gid},
            "cell": {"userEnteredFormat": {
                "backgroundColor": BLANC,
                "textFormat": {"fontFamily": "Inter", "fontSize": 10}}},
            "fields": "userEnteredFormat(backgroundColor,textFormat.fontFamily,textFormat.fontSize)"}})
        req.append({"updateSheetProperties": {
            "properties": {"sheetId": gid, "gridProperties": {"hideGridlines": True}},
            "fields": "gridProperties.hideGridlines"}})

    # mises en avant spécifiques
    if u"Paramètres" in ids:
        req.append({"repeatCell": {
            "range": {"sheetId": ids[u"Paramètres"], "startColumnIndex": 2, "endColumnIndex": 3,
                      "startRowIndex": 4},
            "cell": {"userEnteredFormat": {"backgroundColor": CREME,
                                           "textFormat": {"bold": True},
                                           "horizontalAlignment": "RIGHT"}},
            "fields": "userEnteredFormat(backgroundColor,textFormat.bold,horizontalAlignment)"}})
    if u"MRR mensuel" in ids:
        gid = ids[u"MRR mensuel"]
        req.append({"updateSheetProperties": {
            "properties": {"sheetId": gid,
                           "gridProperties": {"frozenRowCount": 3, "frozenColumnCount": 1}},
            "fields": "gridProperties(frozenRowCount,frozenColumnCount)"}})
        req.append({"repeatCell": {
            "range": {"sheetId": gid, "startRowIndex": 2, "endRowIndex": 3},
            "cell": {"userEnteredFormat": {
                "backgroundColor": BLEU,
                "textFormat": {"bold": True, "foregroundColor": BLANC},
                "horizontalAlignment": "CENTER", "wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,wrapStrategy)"}})
    if u"Synthèse annuelle" in ids:
        gid = ids[u"Synthèse annuelle"]
        req.append({"updateSheetProperties": {
            "properties": {"sheetId": gid, "gridProperties": {"frozenRowCount": 4}},
            "fields": "gridProperties.frozenRowCount"}})
        req.append({"repeatCell": {
            "range": {"sheetId": gid, "startRowIndex": 3, "endRowIndex": 4},
            "cell": {"userEnteredFormat": {
                "backgroundColor": BLEU,
                "textFormat": {"bold": True, "foregroundColor": BLANC},
                "horizontalAlignment": "CENTER", "wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,wrapStrategy)"}})

    for i in range(0, len(req), 200):
        sheets_api.spreadsheets().batchUpdate(
            spreadsheetId=sid, body={"requests": req[i:i + 200]}).execute()
    print(u"Mise en forme appliquée (%d requêtes)" % len(req))


# --------------------------------------------------------------------------- #
def _plage(gid, c0, c1, r0, r1):
    return {"sheetId": gid, "startColumnIndex": c0, "endColumnIndex": c1,
            "startRowIndex": r0, "endRowIndex": r1}


def _serie(gid, col, r0, r1, couleur=None, type_=None):
    s = {"series": {"sourceRange": {"sources": [_plage(gid, col, col + 1, r0, r1)]}},
         "targetAxis": "LEFT_AXIS"}
    if couleur:
        s["color"] = couleur
    if type_:
        s["type"] = type_
    return s


def ajouter_graphiques(sheets_api, sid, ids, onglets):
    """Graphiques Google natifs, ancrés sur les feuilles concernées."""
    par_nom = {o["nom"]: o for o in onglets}
    req = []

    # ---- MRR mensuel : lignes de données 4..40 (index 3..40), en-tête ligne 3 (index 2)
    if u"MRR mensuel" in ids:
        gid = ids[u"MRR mensuel"]
        o = par_nom[u"MRR mensuel"]
        r_head, r0, r1 = 2, 2, o["nlignes"] - 1   # on inclut l'en-tête pour les titres de série
        dom = {"domain": {"sourceRange": {"sources": [_plage(gid, 0, 1, r0, r1)]}}}

        def ancrer(spec, col, row):
            return {"addChart": {"chart": {
                "spec": spec,
                "position": {"overlayPosition": {"anchorCell": {
                    "sheetId": gid, "rowIndex": row, "columnIndex": col},
                    "widthPixels": 620, "heightPixels": 300}}}}}

        req.append(ancrer({
            "title": u"MRR et chiffre d'affaires mensuels",
            "basicChart": {"chartType": "LINE", "legendPosition": "BOTTOM_LEGEND",
                           "headerCount": 1, "domains": [dom],
                           "series": [_serie(gid, 5, r0, r1, {"red": 0, "green": 0.72, "blue": 0.66}),
                                      _serie(gid, 6, r0, r1, {"red": 0, "green": 0.443, "blue": 0.89})],
                           "axis": [{"position": "LEFT_AXIS", "title": u"€ par mois"}]}},
            19, 2))
        req.append(ancrer({
            "title": u"Clients actifs",
            "basicChart": {"chartType": "AREA", "legendPosition": "BOTTOM_LEGEND",
                           "headerCount": 1, "domains": [dom],
                           "series": [_serie(gid, 3, r0, r1, {"red": 0.37, "green": 0.36, "blue": 0.90})]}},
            19, 18))
        req.append(ancrer({
            "title": u"Composition du chiffre d'affaires",
            "basicChart": {"chartType": "COLUMN", "stackedType": "STACKED",
                           "legendPosition": "BOTTOM_LEGEND", "headerCount": 1, "domains": [dom],
                           "series": [_serie(gid, 4, r0, r1, {"red": 0, "green": 0.443, "blue": 0.89}),
                                      _serie(gid, 5, r0, r1, {"red": 0, "green": 0.72, "blue": 0.66})]}},
            19, 34))
        req.append(ancrer({
            "title": u"Revenu total net et épargne cumulée",
            "basicChart": {"chartType": "LINE", "legendPosition": "BOTTOM_LEGEND",
                           "headerCount": 1, "domains": [dom],
                           "series": [_serie(gid, 11, r0, r1, {"red": 0, "green": 0.443, "blue": 0.89}),
                                      _serie(gid, 14, r0, r1, {"red": 0.20, "green": 0.78, "blue": 0.35})]}},
            19, 50))

    # ---- Synthèse annuelle : en-tête index 3, données 4..7
    if u"Synthèse annuelle" in ids:
        gid = ids[u"Synthèse annuelle"]
        r0, r1 = 3, 8
        dom = {"domain": {"sourceRange": {"sources": [_plage(gid, 0, 1, r0, r1)]}}}
        req.append({"addChart": {"chart": {
            "spec": {"title": u"CA annuel : déploiement contre abonnements",
                     "basicChart": {"chartType": "COLUMN", "stackedType": "STACKED",
                                    "legendPosition": "BOTTOM_LEGEND", "headerCount": 1,
                                    "domains": [dom],
                                    "series": [_serie(gid, 2, r0, r1, {"red": 0, "green": 0.443, "blue": 0.89}),
                                               _serie(gid, 3, r0, r1, {"red": 0, "green": 0.72, "blue": 0.66})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 11, "columnIndex": 0},
                "widthPixels": 560, "heightPixels": 320}}}}})
        req.append({"addChart": {"chart": {
            "spec": {"title": u"Revenu total net par exercice",
                     "basicChart": {"chartType": "COLUMN", "stackedType": "STACKED",
                                    "legendPosition": "BOTTOM_LEGEND", "headerCount": 1,
                                    "domains": [dom],
                                    "series": [_serie(gid, 7, r0, r1, {"red": 0.56, "green": 0.56, "blue": 0.58}),
                                               _serie(gid, 6, r0, r1, {"red": 0, "green": 0.72, "blue": 0.66})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 11, "columnIndex": 9},
                "widthPixels": 560, "heightPixels": 320}}}}})
        req.append({"addChart": {"chart": {
            "spec": {"title": u"Épargne dégagée par exercice",
                     "basicChart": {"chartType": "COLUMN", "legendPosition": "NO_LEGEND",
                                    "headerCount": 1, "domains": [dom],
                                    "series": [_serie(gid, 10, r0, r1, {"red": 0.20, "green": 0.78, "blue": 0.35})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 29, "columnIndex": 0},
                "widthPixels": 560, "heightPixels": 300}}}}})
        req.append({"addChart": {"chart": {
            "spec": {"title": u"Rendement net par heure travaillée",
                     "basicChart": {"chartType": "COLUMN", "legendPosition": "NO_LEGEND",
                                    "headerCount": 1, "domains": [dom],
                                    "series": [_serie(gid, 14, r0, r1, {"red": 1.0, "green": 0.62, "blue": 0.04})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 29, "columnIndex": 9},
                "widthPixels": 560, "heightPixels": 300}}}}})

    # ---- Scénarios : en-tête index 4, données 5..9
    if u"Scénarios" in ids:
        gid = ids[u"Scénarios"]
        r0, r1 = 4, 10
        dom = {"domain": {"sourceRange": {"sources": [_plage(gid, 1, 2, r0, r1)]}}}
        req.append({"addChart": {"chart": {
            "spec": {"title": u"MRR au 37ᵉ mois selon le scénario",
                     "basicChart": {"chartType": "COLUMN", "legendPosition": "NO_LEGEND",
                                    "headerCount": 1, "domains": [dom],
                                    "series": [_serie(gid, 7, r0, r1, {"red": 0, "green": 0.72, "blue": 0.66})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 12, "columnIndex": 1},
                "widthPixels": 560, "heightPixels": 320}}}}})
        req.append({"addChart": {"chart": {
            "spec": {"title": u"Net consultant moyen par mois",
                     "basicChart": {"chartType": "COLUMN", "legendPosition": "NO_LEGEND",
                                    "headerCount": 1, "domains": [dom],
                                    "series": [_serie(gid, 8, r0, r1, {"red": 0, "green": 0.443, "blue": 0.89})]}},
            "position": {"overlayPosition": {"anchorCell": {
                "sheetId": gid, "rowIndex": 12, "columnIndex": 10},
                "widthPixels": 560, "heightPixels": 320}}}}})

    if req:
        sheets_api.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": req}).execute()
    print(u"Graphiques natifs ajoutés : %d" % len(req))


# --------------------------------------------------------------------------- #
def main():
    forcer_nouveau = "--new" in sys.argv
    drive, sheets_api = services()
    onglets = lire_classeur()
    print(u"Classeur lu : %d onglets" % len(onglets))

    folder_id = dossier_cible(drive)
    sid, neuf = classeur_cible(drive, sheets_api, folder_id, forcer_nouveau)
    ids = preparer_onglets(sheets_api, sid, onglets, neuf)
    pousser_valeurs(sheets_api, sid, onglets)
    mettre_en_forme(sheets_api, sid, ids, onglets)
    ajouter_graphiques(sheets_api, sid, ids, onglets)

    url = "https://docs.google.com/spreadsheets/d/%s/edit" % sid
    print()
    print(u"✓ %s" % TITRE)
    print(u"  dossier : %s" % DOSSIER)
    print(u"  URL     : %s" % url)
    return 0


if __name__ == "__main__":
    sys.exit(main())
